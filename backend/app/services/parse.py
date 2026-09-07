from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path

from openpyxl import load_workbook

from app.models import ParsedTable


def _normalize_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        # Avoid scientific notation surprises for whole floats
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return format(value, "f").rstrip("0").rstrip(".") if "." in format(value, "f") else str(value)
    text = str(value)
    return text


def parse_bytes(data: bytes, filename: str) -> ParsedTable:
    name = filename.lower()
    size = len(data)
    if name.endswith(".csv"):
        return _parse_delimited(data, filename, size, "csv", ",")
    if name.endswith(".tsv"):
        return _parse_delimited(data, filename, size, "tsv", "\t")
    if name.endswith(".json"):
        return _parse_json(data, filename, size)
    if name.endswith(".xlsx") or name.endswith(".xls"):
        return _parse_xlsx(data, filename, size)
    raise ValueError("Unsupported file type. Use CSV, TSV, JSON, or Excel (.xlsx).")


def parse_path(path: Path) -> ParsedTable:
    return parse_bytes(path.read_bytes(), path.name)


def _parse_delimited(
    data: bytes, filename: str, size: int, fmt: str, delimiter: str
) -> ParsedTable:
    text = data.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = list(reader)
    if not rows:
        raise ValueError("File is empty.")
    headers = [h.strip() or f"column_{i+1}" for i, h in enumerate(rows[0])]
    body: list[list[str]] = []
    for row in rows[1:]:
        if not any(cell.strip() for cell in row):
            continue
        padded = list(row) + [""] * max(0, len(headers) - len(row))
        body.append([_normalize_cell(c) for c in padded[: len(headers)]])
    return ParsedTable(
        headers=headers, rows=body, filename=filename, size_bytes=size, format=fmt  # type: ignore[arg-type]
    )


def _parse_json(data: bytes, filename: str, size: int) -> ParsedTable:
    payload = json.loads(data.decode("utf-8-sig"))
    if isinstance(payload, dict):
        payload = [payload]
    if not isinstance(payload, list) or not payload:
        raise ValueError("JSON must be a non-empty array of objects.")
    headers: list[str] = []
    seen: set[str] = set()
    for item in payload:
        if not isinstance(item, dict):
            raise ValueError("JSON array items must be objects.")
        for key in item.keys():
            if key not in seen:
                seen.add(key)
                headers.append(str(key))
    rows: list[list[str]] = []
    for item in payload:
        row = []
        for h in headers:
            if h not in item or item[h] is None:
                row.append("")
            else:
                # Preserve JSON string identity: quoted numbers stay as their string value
                val = item[h]
                if isinstance(val, str):
                    row.append(val)
                elif isinstance(val, bool):
                    row.append("true" if val else "false")
                elif isinstance(val, int):
                    row.append(str(val))
                elif isinstance(val, float):
                    row.append(repr(val) if "e" in repr(val).lower() else str(val))
                else:
                    row.append(json.dumps(val))
        rows.append(row)
    return ParsedTable(
        headers=headers, rows=rows, filename=filename, size_bytes=size, format="json"
    )


def _parse_xlsx(data: bytes, filename: str, size: int) -> ParsedTable:
    wb = load_workbook(filename=io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.active
    matrix: list[list[str]] = []
    for excel_row in ws.iter_rows(values_only=True):
        if excel_row is None:
            continue
        values = [_excel_cell(c) for c in excel_row]
        if not any(v.strip() for v in values):
            continue
        matrix.append(values)
    if not matrix:
        raise ValueError("Excel sheet is empty.")
    width = max(len(r) for r in matrix)
    headers = [
        (matrix[0][i].strip() if i < len(matrix[0]) and matrix[0][i].strip() else f"column_{i+1}")
        for i in range(width)
    ]
    body: list[list[str]] = []
    for row in matrix[1:]:
        padded = row + [""] * (width - len(row))
        body.append(padded[:width])
    return ParsedTable(
        headers=headers, rows=body, filename=filename, size_bytes=size, format="xlsx"
    )


def _excel_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    # Numbers from Excel lose leading zeros at the Excel layer; keep textual form we receive
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        text = format(value, ".15f").rstrip("0").rstrip(".")
        return text
    return str(value)


_SAFE_TABLE = re.compile(r"[^a-zA-Z0-9_]+")


def suggest_table_name(filename: str) -> str:
    stem = Path(filename).stem
    cleaned = _SAFE_TABLE.sub("_", stem).strip("_").lower()
    return cleaned or "imported_table"
